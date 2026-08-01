from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "RP_APP_STUDENT_APP_BENCHMARK_AND_INTEGRATION_REPORT.md"
OUTPUT = ROOT / "docs" / "RP_APP_STUDENT_APP_BENCHMARK_AND_INTEGRATION_REPORT.docx"
ONEDRIVE_OUTPUT = Path(
    r"C:\Users\정우현\OneDrive\Desktop\RePERFORMANCE\회사문서\RP 홈페이지"
) / OUTPUT.name

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
GREEN = "0B5A4F"
DEEP_GREEN = "173E36"
SAGE = "DDEBE6"
LIGHT_SAGE = "F3F7F5"
WARM = "F5F1E8"
GOLD = "B77B38"
INK = "23302D"
MUTED = "64706C"
BORDER = "CBD8D3"
WHITE = "FFFFFF"


def set_run_font(
    run,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    font_name: str = "Calibri",
):
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths: list[int]):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table width must total {CONTENT_WIDTH_DXA}: {widths}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GREEN)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Malgun Gothic")
    r_pr.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, *, size=10.5, color=INK, bold=False):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size, color=color, bold=bold)
        token = match.group(0)
        if token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size - 0.3, color=DEEP_GREEN, bold=True, font_name="Consolas")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color, bold=bold)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, GREEN, 18, 10),
        "Heading 2": (13, GREEN, 14, 7),
        "Heading 3": (11.5, DEEP_GREEN, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15


def add_numbering(doc: Document, num_format: str, marker: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), marker)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.extend([tabs, indent])
    level.extend([start, fmt, lvl_text, lvl_jc, p_pr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def set_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_ordered_item(paragraph, number: str, text: str):
    paragraph.paragraph_format.left_indent = Inches(0.36)
    paragraph.paragraph_format.first_line_indent = Inches(-0.36)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    prefix = paragraph.add_run(f"{number}. ")
    set_run_font(prefix, size=10.5, color=INK)
    add_inline(paragraph, text, size=10.5, color=INK)


def set_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("RP APP 연구 보고서  |  ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.extend([begin, instr, separate, value, end])
    paragraph._p.append(field_run)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("RePERFORMANCE  |  STUDENT EXPERIENCE RESEARCH")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    set_page_field(footer.paragraphs[0])


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(112)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("PRODUCT RESEARCH · 2026")
    set_run_font(run, size=10, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("RP APP 수험생 앱 벤치마크")
    set_run_font(run, size=29, color=DEEP_GREEN, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("통합 기능 설계 보고서")
    set_run_font(run, size=20, color=GREEN, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(62)
    add_inline(
        p,
        "공부·실기·회복·입시·소통을 한 흐름으로 연결하는 학생 운영체제 설계",
        size=12,
        color=MUTED,
    )

    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=SAGE, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_SAGE)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    add_inline(
        p,
        "핵심 제안  |  RP가 모든 콘텐츠를 복제하기보다 계획→실행→기록→피드백→조정을 연결한다.",
        size=11,
        color=DEEP_GREEN,
        bold=True,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("RePERFORMANCE")
    set_run_font(run, size=12, color=DEEP_GREEN, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("기준일 2026-08-02  |  내부 제품·기술 기획 자료")
    set_run_font(run, size=9.5, color=MUTED)


def extract_main_headings(lines: list[str]) -> list[str]:
    return [line[3:].strip() for line in lines if line.startswith("## ")]


def add_contents(doc: Document, headings: list[str]):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("CONTENTS")
    set_run_font(run, size=9.5, color=GOLD, bold=True)

    h = doc.add_paragraph("목차", style="Heading 1")
    h.paragraph_format.space_before = Pt(0)
    h.paragraph_format.space_after = Pt(14)

    for number, heading in enumerate(headings, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        clean = re.sub(r"^\d+\.\s*", "", heading)
        add_ordered_item(p, str(number), clean)


def table_widths(column_count: int) -> list[int]:
    if column_count == 2:
        return [2520, 6840]
    if column_count == 3:
        return [1980, 3330, 4050]
    if column_count == 4:
        return [1440, 2340, 2880, 2700]
    base = CONTENT_WIDTH_DXA // column_count
    result = [base] * column_count
    result[-1] += CONTENT_WIDTH_DXA - sum(result)
    return result


def add_markdown_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    column_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=column_count)
    set_table_geometry(table, table_widths(column_count))
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    table.paragraph_format = None

    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, SAGE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_SAGE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            add_inline(
                p,
                value,
                size=9.1,
                color=DEEP_GREEN if row_index == 0 else INK,
                bold=row_index == 0,
            )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw_rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        raw_rows.append(parts)
        index += 1
    if len(raw_rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in raw_rows[1]):
        raw_rows.pop(1)
    return raw_rows, index


PAGE_BREAK_HEADINGS = {
    "3. 수험생이 사용하는 앱 영역",
    "5. RP APP 목표 기능 구조",
    "6. 신규 기능 상세안",
    "7. RP AI 도우미 설계",
    "10. 구현 우선순위",
    "15. 조사 출처",
}


def add_body(doc: Document, lines: list[str], bullet_num_id: int):
    index = 0
    first_title_skipped = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("# ") and not first_title_skipped:
            first_title_skipped = True
            index += 1
            continue

        if stripped.startswith("## "):
            text = stripped[3:].strip()
            if text in PAGE_BREAK_HEADINGS:
                doc.add_page_break()
            p = doc.add_paragraph(text, style="Heading 1")
            p.paragraph_format.keep_with_next = True
            index += 1
            continue

        if stripped.startswith("### "):
            p = doc.add_paragraph(stripped[4:].strip(), style="Heading 2")
            p.paragraph_format.keep_with_next = True
            index += 1
            continue

        if stripped.startswith("#### "):
            p = doc.add_paragraph(stripped[5:].strip(), style="Heading 3")
            p.paragraph_format.keep_with_next = True
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue

        if stripped.startswith("> "):
            table = doc.add_table(rows=1, cols=1)
            set_table_geometry(table, [CONTENT_WIDTH_DXA])
            set_table_borders(table, color=SAGE, size="4")
            cell = table.cell(0, 0)
            set_cell_shading(cell, WARM)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, stripped[2:].strip(), size=11, color=DEEP_GREEN, bold=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            index += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph()
            set_numbering(p, bullet_num_id)
            add_inline(p, stripped[2:].strip(), size=10.5, color=INK)
            index += 1
            continue

        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            p = doc.add_paragraph()
            add_ordered_item(p, ordered.group(1), ordered.group(2))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                index += 1
                break
            if (
                candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith("> ")
                or candidate.startswith("- ")
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines), size=10.5, color=INK)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])

    bullet_num_id = add_numbering(doc, "bullet", "•")

    add_cover(doc)
    doc.add_page_break()
    add_contents(doc, extract_main_headings(lines))
    doc.add_page_break()
    add_body(doc, lines, bullet_num_id)

    core = doc.core_properties
    core.title = "RP APP 수험생 앱 벤치마크 및 통합 기능 설계 보고서"
    core.subject = "수험생 앱 기능 조사와 RP APP 통합 로드맵"
    core.author = "RePERFORMANCE"
    core.keywords = "RP APP, 체대입시, 수험생, 제품 기획, 벤치마크"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    ONEDRIVE_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, ONEDRIVE_OUTPUT)
    print(OUTPUT)
    print(ONEDRIVE_OUTPUT)


if __name__ == "__main__":
    build()
